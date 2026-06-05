import argparse
import csv
import hashlib
import json
import re
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib import parse as urlparse
from urllib import request as urlrequest

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


FINAL_GT_SHEET_INDEX = 3
DEFAULT_GT_WORKBOOK = r"D:\data-annotation\final_gt_202506.xlsx"


def clean_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def norm_text(value: Any) -> str:
    return re.sub(r"[\s,，、/\\_：:;；()（）\[\]【】<>《》-]+", "", clean_text(value)).lower()


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def safe_filename(value: str) -> str:
    cleaned = clean_text(value) or datetime.now().strftime("%Y%m%d-%H%M%S")
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", cleaned).strip("_") or "proposal"


def download_proposal(backend: str, project_id: int, source_run_key: str, output_dir: Path) -> Path:
    backend = backend.rstrip("/")
    query = ""
    if clean_text(source_run_key):
        query = "?" + urlparse.urlencode({"sourceRunKey": source_run_key})
    url = f"{backend}/api/projects/{project_id}/problem-gt/export.json{query}"
    req = urlrequest.Request(url, headers={"Accept": "application/json"}, method="GET")
    with urlrequest.urlopen(req, timeout=300) as response:
        body = response.read()
    output_dir.mkdir(parents=True, exist_ok=True)
    proposal_path = output_dir / f"problem_gt_proposal_{safe_filename(source_run_key)}.json"
    proposal_path.write_bytes(body)
    return proposal_path


def write_csv(path: Path, rows: list[dict[str, Any]], headers: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def safe_sheet(wb, sheet_name: str | None, sheet_index: int) -> Worksheet:
    if sheet_name and sheet_name in wb.sheetnames:
        return wb[sheet_name]
    if sheet_index >= len(wb.worksheets):
        raise RuntimeError(f"sheet index {sheet_index} is out of range: {wb.sheetnames}")
    return wb.worksheets[sheet_index]


def header_map(ws: Worksheet) -> dict[str, int]:
    result: dict[str, int] = {}
    for cell in ws[1]:
        header = clean_text(cell.value)
        if header:
            result[header] = cell.column
    return result


def max_problem_id_number(ws: Worksheet, problem_id_col: int, prefix: str) -> int:
    max_no = 0
    pattern = re.compile(re.escape(prefix) + r"_(\d+)$")
    for row in range(2, ws.max_row + 1):
        value = clean_text(ws.cell(row=row, column=problem_id_col).value)
        match = pattern.search(value)
        if match:
            max_no = max(max_no, int(match.group(1)))
    return max_no


def existing_fingerprints(ws: Worksheet, headers: dict[str, int]) -> set[tuple[str, ...]]:
    def value(row: int, header: str) -> str:
        col = headers.get(header)
        return clean_text(ws.cell(row=row, column=col).value) if col else ""

    result: set[tuple[str, ...]] = set()
    for row in range(2, ws.max_row + 1):
        issue_type = value(row, "问题类型（schema）")
        if not issue_type:
            continue
        result.add(
            (
                norm_text(value(row, "pair-id")),
                norm_text(issue_type),
                norm_text(value(row, "浅度规范化-row_hierarchy") or value(row, "源excel中路径") or value(row, "源pdf中路径")),
                norm_text(value(row, "浅度规范化-column_hierarchy")),
                norm_text(value(row, "源pdf中value")),
                norm_text(value(row, "源excel中value")),
            )
        )
    return result


def proposal_fingerprint(row: dict[str, Any]) -> tuple[str, ...]:
    return (
        norm_text(row.get("note_no")),
        norm_text(row.get("final_issue_type")),
        norm_text(row.get("row_label") or row.get("excel_locator") or row.get("pdf_locator")),
        norm_text(row.get("column_label")),
        norm_text(row.get("pdf_value")),
        norm_text(row.get("excel_value")),
    )


def next_problem_id(prefix: str, number: int) -> str:
    return f"{prefix}_{number:04d}"


def semantic_path(row: dict[str, Any]) -> str:
    parts = [
        row.get("note_name"),
        row.get("row_label"),
        row.get("column_label"),
        row.get("period_label"),
    ]
    return " / ".join(clean_text(part) for part in parts if clean_text(part))


def copy_row_style(ws: Worksheet, template_row: int, target_row: int) -> None:
    ws.row_dimensions[target_row].height = ws.row_dimensions[template_row].height
    for col in range(1, ws.max_column + 1):
        source = ws.cell(row=template_row, column=col)
        target = ws.cell(row=target_row, column=col)
        if source.has_style:
            target._style = copy(source._style)
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy(source.protection)


def set_cell(ws: Worksheet, row_idx: int, headers: dict[str, int], header: str, value: Any) -> None:
    col = headers.get(header)
    if col:
        ws.cell(row=row_idx, column=col, value=value)


def write_gt_row(
    ws: Worksheet,
    headers: dict[str, int],
    target_row: int,
    proposal_row: dict[str, Any],
    problem_id: str,
    source_run_key: str,
) -> dict[str, Any]:
    note_no = clean_text(proposal_row.get("note_no"))
    note_name = clean_text(proposal_row.get("note_name"))
    issue_type = clean_text(proposal_row.get("final_issue_type") or proposal_row.get("suggested_issue_type"))
    pdf_locator = clean_text(proposal_row.get("pdf_locator"))
    excel_locator = clean_text(proposal_row.get("excel_locator"))
    path = semantic_path(proposal_row)
    rationale = clean_text(proposal_row.get("rationale") or proposal_row.get("review_comment") or proposal_row.get("system_reason"))

    values = {
        "联查分组": "问题GT助手-新增",
        "联查状态": "scoreable_gt",
        "是否标黄": "未标明",
        "源pdf中定位": pdf_locator or "empty",
        "源excel中定位": excel_locator or "empty",
        "pair-id": note_no,
        "附注名": note_name,
        "问题项id": problem_id,
        "是否最终GT": "是",
        "问题类型（schema）": issue_type,
        "源pdf中路径": pdf_locator or path or "empty",
        "源excel中路径": excel_locator or path or "empty",
        "源pdf中value": clean_text(proposal_row.get("pdf_value")) or "empty",
        "源excel中value": clean_text(proposal_row.get("excel_value")) or "empty",
        "浅度规范化-note_table_path": note_name,
        "浅度规范化-section_hierarchy": "",
        "浅度规范化-row_hierarchy": clean_text(proposal_row.get("row_label")),
        "浅度规范化-column_hierarchy": clean_text(proposal_row.get("column_label")),
        "联查说明": f"gt-review-assistant proposal add; run={source_run_key}; candidate={clean_text(proposal_row.get('candidate_id'))}",
        "审核说明": rationale,
    }
    for header, value in values.items():
        set_cell(ws, target_row, headers, header, value)
    return {
        "problem_id": problem_id,
        "candidate_id": clean_text(proposal_row.get("candidate_id")),
        "candidate_key": clean_text(proposal_row.get("candidate_key")),
        "decision": clean_text(proposal_row.get("decision")),
        "note_no": note_no,
        "note_name": note_name,
        "issue_type": issue_type,
        "row_label": clean_text(proposal_row.get("row_label")),
        "column_label": clean_text(proposal_row.get("column_label")),
        "pdf_value": clean_text(proposal_row.get("pdf_value")),
        "excel_value": clean_text(proposal_row.get("excel_value")),
        "pdf_locator": pdf_locator,
        "excel_locator": excel_locator,
        "rationale": rationale,
    }


def add_final_rows(
    workbook_path: Path,
    proposal: dict[str, Any],
    output_workbook: Path,
    sheet_name: str,
    sheet_index: int,
    problem_id_prefix: str,
    allow_duplicates: bool,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    wb = load_workbook(workbook_path)
    ws = safe_sheet(wb, sheet_name, sheet_index)
    headers = header_map(ws)
    missing_headers = [
        header
        for header in ["问题项id", "问题类型（schema）", "pair-id", "附注名", "是否最终GT"]
        if header not in headers
    ]
    if missing_headers:
        raise RuntimeError(f"final GT sheet missing required headers: {missing_headers}")

    source_run_key = clean_text(proposal.get("source_run_key"))
    adds = list(proposal.get("adds", []))
    existing = existing_fingerprints(ws, headers)
    next_no = max_problem_id_number(ws, headers["问题项id"], problem_id_prefix) + 1
    template_row = max(2, ws.max_row)
    applied: list[dict[str, Any]] = []
    skipped_duplicates: list[dict[str, Any]] = []

    for proposal_row in adds:
        fingerprint = proposal_fingerprint(proposal_row)
        if not allow_duplicates and fingerprint in existing:
            skipped_duplicates.append(proposal_row)
            continue
        target_row = ws.max_row + 1
        copy_row_style(ws, template_row, target_row)
        problem_id = next_problem_id(problem_id_prefix, next_no)
        next_no += 1
        applied_row = write_gt_row(ws, headers, target_row, proposal_row, problem_id, source_run_key)
        applied.append(applied_row)
        existing.add(fingerprint)

    output_workbook.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_workbook)
    meta = {
        "sheet": ws.title,
        "original_data_rows": max(template_row - 1, 0),
        "new_data_rows": max(ws.max_row - 1, 0),
        "proposal_add_rows": len(adds),
        "applied_rows": len(applied),
        "skipped_duplicate_rows": len(skipped_duplicates),
    }
    return applied, skipped_duplicates, meta


def export_headers() -> list[str]:
    return [
        "problem_id",
        "candidate_id",
        "candidate_key",
        "decision",
        "note_no",
        "note_name",
        "issue_type",
        "row_label",
        "column_label",
        "pdf_value",
        "excel_value",
        "pdf_locator",
        "excel_locator",
        "rationale",
    ]


def write_summary(
    output_dir: Path,
    summary: dict[str, Any],
    applied: list[dict[str, Any]],
    skipped_duplicates: list[dict[str, Any]],
    proposal: dict[str, Any],
) -> None:
    excluded = list(proposal.get("excluded", []))
    unresolved = list(proposal.get("unresolved", []))
    all_candidates = list(proposal.get("all_candidates", []))
    skipped_headers = sorted({k for row in skipped_duplicates for k in row.keys()}) or ["candidate_id", "candidate_key", "decision", "reason"]
    excluded_headers = sorted({k for row in excluded for k in row.keys()}) or ["candidate_id", "candidate_key", "decision", "reason"]
    unresolved_headers = sorted({k for row in unresolved for k in row.keys()}) or ["candidate_id", "candidate_key", "decision", "reason"]
    write_csv(output_dir / "applied_add_rows.csv", applied, export_headers())
    write_csv(output_dir / "skipped_duplicate_rows.csv", skipped_duplicates, skipped_headers)
    write_csv(output_dir / "excluded_not_applied.csv", excluded, excluded_headers)
    write_csv(output_dir / "unresolved_not_applied.csv", unresolved, unresolved_headers)
    summary_path = output_dir / "apply_problem_gt_proposal_summary.json"
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    md = f"""# 问题GT Proposal 安全写回结果

- 状态：只生成新 workbook；没有覆盖原始 final_gt。
- 原始 workbook：`{summary['input_workbook']}`
- 新 workbook：`{summary['output_workbook']}`
- 原始 SHA256：`{summary['input_sha256']}`
- 新文件 SHA256：`{summary['output_sha256']}`
- proposal run：`{summary['source_run_key']}`

## 数量

- proposal 全部候选：{len(all_candidates)}
- proposal 建议补入：{summary['proposal_add_rows']}
- 实际写入新增：{summary['applied_rows']}
- 因疑似已存在而跳过：{summary['skipped_duplicate_rows']}
- 排除/重复/表级合并：{len(excluded)}（只记录，不删除原 GT）
- 未完成/待复核：{len(unresolved)}（只记录，不写入）
- 原 final 数据行数：{summary['original_data_rows']}
- 新 final 数据行数：{summary['new_data_rows']}

## 安全边界

- 本脚本默认只追加 `proposal.adds`，不会删除、改写、覆盖原 workbook。
- `excluded_not_applied.csv` 只是留痕，不代表已经从 final GT 删除。
- `unresolved_not_applied.csv` 需要继续人工处理，不能直接冻结。
"""
    (output_dir / "apply_problem_gt_proposal_summary.md").write_text(md, encoding="utf-8")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="安全应用问题GT proposal：只生成新版 workbook，不覆盖原文件")
    parser.add_argument("--gt-workbook", default=DEFAULT_GT_WORKBOOK, help="原始 final_gt workbook")
    parser.add_argument("--proposal-json", default="", help="后端导出的 problem_gt proposal JSON；不填时可用 --backend 自动下载")
    parser.add_argument("--backend", default="", help="后端地址，例如 http://localhost:8080；配合 --project-id 自动下载 proposal")
    parser.add_argument("--project-id", type=int, default=0, help="项目 ID；配合 --backend 自动下载 proposal")
    parser.add_argument("--source-run-key", default="", help="问题GT候选 run key；不填则由后端使用最新 run")
    parser.add_argument("--output-dir", required=True, help="输出目录")
    parser.add_argument("--output-workbook", default="", help="可选：指定新版 workbook 路径")
    parser.add_argument("--sheet-name", default="最终GT")
    parser.add_argument("--sheet-index", type=int, default=FINAL_GT_SHEET_INDEX)
    parser.add_argument("--problem-id-prefix", default="FGT202506")
    parser.add_argument("--allow-duplicates", action="store_true", help="默认会跳过疑似已存在项；打开后允许重复追加")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    workbook_path = Path(args.gt_workbook)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    if not workbook_path.exists():
        raise FileNotFoundError(workbook_path)
    if args.proposal_json:
        proposal_path = Path(args.proposal_json)
        if not proposal_path.exists():
            raise FileNotFoundError(proposal_path)
    else:
        if not args.backend or not args.project_id:
            raise RuntimeError("either --proposal-json or both --backend and --project-id are required")
        proposal_path = download_proposal(args.backend, args.project_id, args.source_run_key, output_dir)

    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    output_workbook = Path(args.output_workbook) if args.output_workbook else output_dir / f"{workbook_path.stem}_problem_gt_applied_{timestamp}.xlsx"
    proposal = read_json(proposal_path)
    input_sha = file_sha256(workbook_path)
    applied, skipped_duplicates, meta = add_final_rows(
        workbook_path=workbook_path,
        proposal=proposal,
        output_workbook=output_workbook,
        sheet_name=args.sheet_name,
        sheet_index=args.sheet_index,
        problem_id_prefix=args.problem_id_prefix,
        allow_duplicates=args.allow_duplicates,
    )
    output_sha = file_sha256(output_workbook)
    summary = {
        "status": "new_workbook_generated_original_not_modified",
        "input_workbook": str(workbook_path),
        "output_workbook": str(output_workbook),
        "proposal_json": str(proposal_path),
        "input_sha256": input_sha,
        "output_sha256": output_sha,
        "source_run_key": clean_text(proposal.get("source_run_key")),
        **meta,
        "excluded_not_applied_rows": len(proposal.get("excluded", [])),
        "unresolved_not_applied_rows": len(proposal.get("unresolved", [])),
        "allow_duplicates": args.allow_duplicates,
    }
    write_summary(output_dir, summary, applied, skipped_duplicates, proposal)
    print(f"status={summary['status']}")
    print(f"output_workbook={output_workbook}")
    print(f"applied_rows={len(applied)} skipped_duplicates={len(skipped_duplicates)} unresolved={summary['unresolved_not_applied_rows']}")
    print(f"summary={output_dir / 'apply_problem_gt_proposal_summary.md'}")


if __name__ == "__main__":
    main()
