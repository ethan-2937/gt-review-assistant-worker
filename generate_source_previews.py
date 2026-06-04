from __future__ import annotations

import argparse
import hashlib
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import fitz
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from PIL import Image, ImageDraw, ImageFont


@dataclass(frozen=True)
class PreviewLocator:
    side: str
    locator: str
    level: str


def md5_text(text: str) -> str:
    return hashlib.md5(text.encode("utf-8")).hexdigest()


def read_json(path: str | Path) -> dict:
    with Path(path).open("r", encoding="utf-8") as handle:
        return json.load(handle)


def clean_text(value) -> str:
    return re.sub(r"\s+", " ", "" if value is None else str(value)).strip()


def parse_excel_locator(locator: str) -> tuple[str, str] | None:
    parts = [part.strip() for part in locator.split("!") if part.strip()]
    if len(parts) < 3:
        return None
    cell = parts[-1]
    if not re.fullmatch(r"[A-Za-z]{1,5}\d+", cell):
        return None
    return parts[-2], cell.upper()


def pdf_page_locator(locator: str) -> str:
    parts = [part.strip() for part in locator.split("!") if part.strip()]
    if len(parts) >= 2 and re.fullmatch(r"p\d+", parts[1], flags=re.IGNORECASE):
        return "!".join(parts[:2])
    match = re.search(r"p(\d+)", locator, flags=re.IGNORECASE)
    if match and parts:
        return f"{parts[0]}!p{match.group(1)}"
    return locator


def parse_pdf_page(locator: str) -> int | None:
    match = re.search(r"(?:^|!)p(\d+)(?:!|$)", locator, flags=re.IGNORECASE)
    if not match:
        return None
    return max(int(match.group(1)), 1)


def iter_structure_locators(payload: dict, *, note_no: str, levels: set[str]) -> Iterable[PreviewLocator]:
    side = payload.get("side", "")
    for note in payload.get("notes", []):
        if note_no and str(note.get("noteNo", "")) != str(note_no):
            continue
        if "note" in levels and note.get("sourceLocator"):
            yield PreviewLocator(side, note["sourceLocator"], "note")
        for table in note.get("tables", []):
            if "table" in levels and table.get("sourceLocator"):
                yield PreviewLocator(side, table["sourceLocator"], "table")
            if "row" in levels:
                for row in table.get("rows", []):
                    if row.get("sourceLocator"):
                        yield PreviewLocator(side, row["sourceLocator"], "row")
            if "column" in levels:
                for col in table.get("columns", []):
                    if col.get("sourceLocator"):
                        yield PreviewLocator(side, col["sourceLocator"], "column")
            if "cell" in levels:
                for cell in table.get("cells", []):
                    if cell.get("sourceLocator"):
                        yield PreviewLocator(side, cell["sourceLocator"], "cell")


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        Path("C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/simsun.ttc"),
    ]
    for candidate in candidates:
        if candidate.exists():
            try:
                return ImageFont.truetype(str(candidate), size=size)
            except Exception:
                continue
    return ImageFont.load_default()


def cell_fill(cell) -> tuple[int, int, int]:
    color = getattr(cell.fill.fgColor, "rgb", None)
    if color and isinstance(color, str) and len(color) in (6, 8):
        rgb = color[-6:]
        try:
            return tuple(int(rgb[idx:idx + 2], 16) for idx in (0, 2, 4))
        except ValueError:
            pass
    return 255, 255, 255


def text_color_for(bg: tuple[int, int, int]) -> tuple[int, int, int]:
    brightness = (bg[0] * 299 + bg[1] * 587 + bg[2] * 114) / 1000
    return (30, 30, 30) if brightness > 145 else (255, 255, 255)


def render_excel_preview(excel_path: Path, locator: str, output_path: Path) -> bool:
    parsed = parse_excel_locator(locator)
    if not parsed:
        return False
    sheet_name, cell_ref = parsed
    wb = load_workbook(excel_path, data_only=True, read_only=False)
    try:
        if sheet_name not in wb.sheetnames:
            return False
        ws = wb[sheet_name]
        match = re.fullmatch(r"([A-Z]+)(\d+)", cell_ref)
        if not match:
            return False
        target_col = column_index_from_string(match.group(1))
        target_row = int(match.group(2))
        min_row = max(1, target_row - 5)
        max_row = min(ws.max_row, target_row + 5)
        min_col = max(1, target_col - 3)
        max_col = min(ws.max_column, target_col + 5)

        row_h = 34
        col_w = 145
        row_head_w = 54
        title_h = 52
        width = row_head_w + (max_col - min_col + 1) * col_w + 2
        height = title_h + row_h + (max_row - min_row + 1) * row_h + 2
        image = Image.new("RGB", (width, height), (250, 248, 241))
        draw = ImageDraw.Draw(image)
        normal_font = font(14)
        small_font = font(12)
        title_font = font(18, bold=True)

        draw.rectangle([0, 0, width, title_h], fill=(36, 79, 63))
        draw.text((14, 13), f"{excel_path.name} / {sheet_name} / {cell_ref}", fill=(255, 250, 240), font=title_font)

        for col in range(min_col, max_col + 1):
            x = row_head_w + (col - min_col) * col_w
            draw.rectangle([x, title_h, x + col_w, title_h + row_h], fill=(231, 241, 235), outline=(193, 203, 194))
            draw.text((x + 8, title_h + 8), get_column_letter(col), fill=(48, 54, 47), font=small_font)

        for row in range(min_row, max_row + 1):
            y = title_h + row_h + (row - min_row) * row_h
            draw.rectangle([0, y, row_head_w, y + row_h], fill=(231, 241, 235), outline=(193, 203, 194))
            draw.text((8, y + 8), str(row), fill=(48, 54, 47), font=small_font)
            for col in range(min_col, max_col + 1):
                x = row_head_w + (col - min_col) * col_w
                cell = ws.cell(row, col)
                bg = cell_fill(cell)
                if bg == (0, 0, 0):
                    bg = (255, 255, 255)
                draw.rectangle([x, y, x + col_w, y + row_h], fill=bg, outline=(213, 213, 205))
                value = clean_text(cell.value)
                if len(value) > 18:
                    value = value[:18] + "..."
                draw.text((x + 7, y + 8), value, fill=text_color_for(bg), font=normal_font)
                if row == target_row and col == target_col:
                    draw.rectangle([x + 2, y + 2, x + col_w - 2, y + row_h - 2], outline=(196, 85, 45), width=4)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        image.save(output_path)
        return True
    finally:
        wb.close()


def render_pdf_preview(pdf_path: Path, page_locator: str, output_path: Path) -> bool:
    page_no = parse_pdf_page(page_locator)
    if not page_no:
        return False
    doc = fitz.open(pdf_path)
    try:
        page_index = page_no - 1
        if page_index < 0 or page_index >= doc.page_count:
            return False
        page = doc.load_page(page_index)
        zoom = 1.25
        pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        pix.save(str(output_path))
        return True
    finally:
        doc.close()


def generate_previews(
    *,
    project_id: str,
    output_dir: Path,
    excel_json: Path | None,
    pdf_json: Path | None,
    excel_input: Path | None,
    pdf_input: Path | None,
    note_no: str,
    levels: set[str],
    max_items: int,
) -> dict[str, int]:
    stats = {"excel": 0, "pdf": 0, "skipped": 0}

    if excel_json and excel_json.exists():
        payload = read_json(excel_json)
        source_path = excel_input or Path(payload.get("sourceFilePath", ""))
        seen: set[str] = set()
        for item in iter_structure_locators(payload, note_no=note_no, levels=levels):
            parsed = parse_excel_locator(item.locator)
            if not parsed:
                stats["skipped"] += 1
                continue
            key = item.locator
            if key in seen:
                continue
            seen.add(key)
            out = output_dir / project_id / "excel" / f"{md5_text(key)}.png"
            if out.exists():
                continue
            if max_items and stats["excel"] >= max_items:
                break
            stats["excel"] += 1 if render_excel_preview(source_path, key, out) else 0

    if pdf_json and pdf_json.exists():
        payload = read_json(pdf_json)
        source_path = pdf_input or Path(payload.get("sourceFilePath", ""))
        seen: set[str] = set()
        for item in iter_structure_locators(payload, note_no=note_no, levels=levels):
            key = pdf_page_locator(item.locator)
            if key in seen:
                continue
            seen.add(key)
            out = output_dir / project_id / "pdf" / f"{md5_text(key)}.png"
            if out.exists():
                continue
            if max_items and stats["pdf"] >= max_items:
                break
            stats["pdf"] += 1 if render_pdf_preview(source_path, key, out) else 0

    return stats


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate real PDF/Excel preview PNGs for structure locators.")
    parser.add_argument("--project-id", required=True)
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--excel-json", default="")
    parser.add_argument("--pdf-json", default="")
    parser.add_argument("--excel-input", default="")
    parser.add_argument("--pdf-input", default="")
    parser.add_argument("--note-no", default="")
    parser.add_argument("--levels", default="row", help="Comma separated: note,table,row,column,cell")
    parser.add_argument("--max-items", type=int, default=0, help="Per side cap; 0 means no cap")
    args = parser.parse_args()

    levels = {part.strip() for part in args.levels.split(",") if part.strip()}
    stats = generate_previews(
        project_id=str(args.project_id),
        output_dir=Path(args.output_dir),
        excel_json=Path(args.excel_json) if args.excel_json else None,
        pdf_json=Path(args.pdf_json) if args.pdf_json else None,
        excel_input=Path(args.excel_input) if args.excel_input else None,
        pdf_input=Path(args.pdf_input) if args.pdf_input else None,
        note_no=str(args.note_no).strip(),
        levels=levels,
        max_items=args.max_items,
    )
    print(json.dumps(stats, ensure_ascii=False))


if __name__ == "__main__":
    main()
