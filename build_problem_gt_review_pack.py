import argparse
import csv
import hashlib
import html
import json
import re
import zipfile
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib import request as urlrequest
from xml.etree import ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


FINAL_GT_SHEET_INDEX = 3
LOCKED_NOTE_NAMES = {
    1: "货币资金",
    2: "结算备付金",
    3: "融出资金",
    4: "衍生金融工具",
    5: "买入返售金融资产",
    6: "应收款项",
    7: "存出保证金",
    8: "金融投资：交易性金融资产",
    9: "金融投资：其他债权投资",
    10: "金融投资：其他权益工具投资",
    11: "融券业务",
}


def clean_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def clip_text(value: Any, max_chars: int = 2000) -> str:
    text = clean_text(value)
    if max_chars <= 0 or len(text) <= max_chars:
        return text
    return f"{text[:max_chars]}... [truncated {len(text) - max_chars} chars]"


def note_int(value: Any) -> int | None:
    match = re.search(r"\d+", str(value or ""))
    return int(match.group()) if match else None


def norm_text(value: Any) -> str:
    return re.sub(r"[\s,，;；|/／:_：()（）\[\]【】\-]+", "", clean_text(value)).lower()


def norm_num(value: Any) -> str:
    text = clean_text(value)
    if not text:
        return ""
    nums = re.findall(r"-?\d+(?:,\d{3})*(?:\.\d+)?|-?\d+(?:\.\d+)?", text)
    return "|".join(num.replace(",", "") for num in nums)


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def post_json(url: str, payload: dict[str, Any]) -> dict[str, Any]:
    data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    req = urlrequest.Request(url, data=data, headers={"Content-Type": "application/json; charset=utf-8"}, method="POST")
    try:
        with urlrequest.urlopen(req, timeout=300) as response:
            return json.loads(response.read().decode("utf-8"))
    except Exception as exc:
        raise RuntimeError(f"POST failed: {url}: {exc}") from exc


def parse_lock_notes(text: str) -> set[int]:
    result: set[int] = set()
    for part in text.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            left, right = part.split("-", 1)
            result.update(range(int(left), int(right) + 1))
        else:
            result.add(int(part))
    return result


def column_number(cell_ref: str) -> int:
    letters = re.sub(r"[^A-Z]", "", cell_ref.upper())
    result = 0
    for letter in letters:
        result = result * 26 + ord(letter) - ord("A") + 1
    return result


def read_shared_strings(zf: zipfile.ZipFile) -> list[str]:
    if "xl/sharedStrings.xml" not in zf.namelist():
        return []
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    shared: list[str] = []
    for event, elem in ET.iterparse(zf.open("xl/sharedStrings.xml"), events=("end",)):
        if elem.tag == ns + "si":
            shared.append("".join(elem.itertext()))
            elem.clear()
    return shared


def workbook_sheet_path(zf: zipfile.ZipFile, sheet_index: int) -> str:
    main_ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    rel_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    package_rel_ns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    workbook = ET.fromstring(zf.read("xl/workbook.xml"))
    sheets = workbook.find(main_ns + "sheets")
    if sheets is None:
        raise RuntimeError("workbook.xml has no sheets node")
    sheet_nodes = list(sheets.findall(main_ns + "sheet"))
    if sheet_index >= len(sheet_nodes):
        raise RuntimeError(f"final GT sheet index {sheet_index} is out of range")
    rel_id = sheet_nodes[sheet_index].attrib.get(rel_ns + "id")
    rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    for rel in rels.findall(package_rel_ns + "Relationship"):
        if rel.attrib.get("Id") == rel_id:
            target = rel.attrib["Target"].lstrip("/")
            return target if target.startswith("xl/") else f"xl/{target}"
    raise RuntimeError(f"sheet relationship not found: {rel_id}")


def read_cell_value(cell: ET.Element, shared_strings: list[str], ns: str) -> str:
    cell_type = cell.attrib.get("t", "")
    if cell_type == "inlineStr":
        return clean_text("".join(cell.itertext()))
    value = cell.find(ns + "v")
    if value is None or value.text is None:
        return ""
    if cell_type == "s":
        try:
            return clean_text(shared_strings[int(value.text)])
        except (ValueError, IndexError):
            return ""
    return clean_text(value.text)


def read_final_gt(path: Path) -> tuple[list[dict[str, Any]], dict[int, list[dict[str, Any]]]]:
    # The final GT workbook can be slow to load through openpyxl because it carries
    # multiple formatted sheets. A tiny XML reader is enough for this read-only scan.
    wanted_columns = {4, 5, 8, 9, 10, 12, 13, 14, 15, 16, 19, 20}
    rows: list[dict[str, Any]] = []
    by_note: dict[int, list[dict[str, Any]]] = defaultdict(list)
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with zipfile.ZipFile(path) as zf:
        sheet_path = workbook_sheet_path(zf, FINAL_GT_SHEET_INDEX)
        shared_strings = read_shared_strings(zf)
        for event, elem in ET.iterparse(zf.open(sheet_path), events=("end",)):
            if elem.tag != ns + "row":
                continue
            row_idx = int(elem.attrib.get("r", "0") or 0)
            if row_idx < 2:
                elem.clear()
                continue
            cells: dict[int, str] = {}
            for cell in elem.findall(ns + "c"):
                col_idx = column_number(cell.attrib.get("r", ""))
                if col_idx in wanted_columns:
                    cells[col_idx] = read_cell_value(cell, shared_strings, ns)
            note_no = note_int(cells.get(8))
            problem_id = clean_text(cells.get(10))
            issue_type = clean_text(cells.get(12))
            if note_no is None and not problem_id and not issue_type:
                elem.clear()
                continue
            record = {
                "excel_row": row_idx,
                "note": note_no,
                "note_name": clean_text(cells.get(9)),
                "issue_type": issue_type,
                "pdf_locator": clean_text(cells.get(4)),
                "excel_locator": clean_text(cells.get(5)),
                "pdf_path": clean_text(cells.get(13)),
                "excel_path": clean_text(cells.get(14)),
                "pdf_value": clean_text(cells.get(15)),
                "excel_value": clean_text(cells.get(16)),
                "row_path": clean_text(cells.get(19)),
                "column_path": clean_text(cells.get(20)),
                "problem_id": problem_id,
            }
            rows.append(record)
            if note_no is not None:
                by_note[note_no].append(record)
            elem.clear()
    return rows, by_note


def read_manual_decisions(path: Path | None) -> dict[str, dict[str, Any]]:
    if not path or not path.exists():
        return {}
    data = read_json(path)
    return {clean_text(item.get("candidate_id")): item for item in data.get("records", [])}


def infer_issue_type(candidate: dict[str, Any]) -> str:
    kind = clean_text(candidate.get("candidate_kind"))
    if kind == "row_label_not_found_and_values_missing":
        return "row:missing"
    if kind == "same_label_value_missing_or_mismatch":
        return "cell:value_mismatch" if candidate.get("row_label_found_in_pdf") else "row:missing"
    if kind == "pdf_number_not_found_in_excel":
        return "candidate:pdf_extra_value"
    return "candidate:unknown"


def final_coverage_hint(candidate: dict[str, Any], final_by_note: dict[int, list[dict[str, Any]]]) -> tuple[str, str]:
    note = note_int(candidate.get("note"))
    if note is None:
        return "", ""
    current_rows = final_by_note.get(note, [])
    row_label = clean_text(candidate.get("row_label"))
    missing_values = clean_text(candidate.get("missing_values"))
    pdf_value = clean_text(candidate.get("pdf_value"))
    row_key = norm_text(row_label)
    candidate_nums = set(filter(None, norm_num(missing_values or pdf_value).split("|")))
    matched: list[str] = []
    for row in current_rows:
        haystack = " ".join(
            [
                row.get("pdf_path", ""),
                row.get("excel_path", ""),
                row.get("pdf_value", ""),
                row.get("excel_value", ""),
                row.get("problem_id", ""),
            ]
        )
        hay_norm = norm_text(haystack)
        if row_key and row_key in hay_norm:
            matched.append(f"row{row['excel_row']}:{row['issue_type']}:{row.get('excel_value') or row.get('pdf_value')}")
            continue
        final_nums = set(filter(None, norm_num(f"{row.get('pdf_value', '')} {row.get('excel_value', '')}").split("|")))
        if candidate_nums and candidate_nums & final_nums:
            matched.append(f"row{row['excel_row']}:{row['issue_type']}:{row.get('excel_value') or row.get('pdf_value')}")
    if matched:
        return "maybe_covered_by_current_final", "; ".join(matched[:5])
    return "", ""


def classify_excel_candidate(
    candidate: dict[str, Any],
    *,
    locked_notes: set[int],
    final_by_note: dict[int, list[dict[str, Any]]],
    manual_decisions: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    note = note_int(candidate.get("note")) or 0
    candidate_id = clean_text(candidate.get("candidate_id"))
    manual = manual_decisions.get(candidate_id)
    overlap = clean_text(candidate.get("current_gt_overlap"))
    current_hint, current_hint_detail = final_coverage_hint(candidate, final_by_note)
    suggested_issue_type = manual.get("issue_type") if manual else infer_issue_type(candidate)

    if note in locked_notes:
        bucket = "locked_confirmed_skip"
        action = "skip_front11_confirmed"
        risk = "skip"
        reason = "附注 1-11 已由人工确认 100% 正确，本 MVP 不再重复审核。"
    elif manual:
        action = clean_text(manual.get("recommended_action"))
        risk = "manual_reused"
        if action == "add":
            bucket = "manual_confirmed_add_check_final"
            reason = "历史人工已确认应补入；用于检查 final_gt 是否已吸收。"
        elif action.startswith("optional"):
            bucket = "manual_optional_or_dedup"
            reason = "历史人工认为可选/需去重策略，不能直接批量加入。"
        else:
            bucket = "manual_no_add_reused"
            reason = "历史人工已排除或判定被覆盖，本轮不再重复筛。"
    elif overlap:
        bucket = "covered_by_source_gt_index"
        action = "no_add"
        risk = "low"
        reason = "历史候选源标记已有 current_gt_overlap，优先视为已覆盖/重复。"
    elif current_hint:
        bucket = current_hint
        action = "needs_quick_check"
        risk = "medium"
        reason = "当前 final_gt 中找到同附注近似行名或数值，疑似已覆盖，需快速确认是否重复。"
    else:
        priority = int(candidate.get("priority") or 9)
        if priority == 1:
            bucket = "high_confidence_possible_missing"
            action = "review_first"
            risk = "high"
            reason = "未锁定、无历史覆盖、无人工排除、优先级 1，最值得先人工确认。"
        else:
            bucket = "medium_confidence_possible_missing"
            action = "review_after_high"
            risk = "medium"
            reason = "未锁定且未覆盖，但候选源优先级为 2，需要结合原文证据确认。"

    return {
        "candidate_id": candidate_id,
        "note": note,
        "folder": clean_text(candidate.get("folder")),
        "source": "excel_to_pdf",
        "bucket": bucket,
        "risk": risk,
        "recommended_action": action,
        "suggested_issue_type": clean_text(suggested_issue_type),
        "reason": reason,
        "candidate_kind": clean_text(candidate.get("candidate_kind")),
        "priority": candidate.get("priority"),
        "period_bucket": clean_text(candidate.get("period_bucket")),
        "xlsx_file": clean_text(candidate.get("xlsx_file")),
        "pdf_file": clean_text(candidate.get("pdf_file")),
        "sheet": clean_text(candidate.get("sheet")),
        "source_row": candidate.get("source_row"),
        "row_label": clean_text(candidate.get("row_label")),
        "column_label": clean_text(manual.get("column_label") if manual else candidate.get("column_hints")),
        "period": clean_text(manual.get("period") if manual else "2025半年度"),
        "source_value": clean_text(manual.get("source_value") if manual else "empty"),
        "target_value": clean_text(manual.get("target_value") if manual else candidate.get("missing_values")),
        "missing_cells": clean_text(candidate.get("missing_cells")),
        "missing_values": clean_text(candidate.get("missing_values")),
        "matched_value_count": candidate.get("matched_value_count"),
        "missing_value_count": candidate.get("missing_value_count"),
        "row_label_found_in_pdf": candidate.get("row_label_found_in_pdf"),
        "current_gt_overlap": overlap,
        "current_final_hint": current_hint_detail,
        "manual_decision": clean_text(manual.get("review_decision") if manual else ""),
        "manual_rationale": clean_text(manual.get("rationale") if manual else ""),
        "pdf_context": clean_text(candidate.get("pdf_context_near_label")),
        "source_locator": clean_text(manual.get("source_locator") if manual else candidate.get("pdf_file")),
        "target_locator": clean_text(
            manual.get("target_locator")
            if manual
            else f"{candidate.get('xlsx_file', '')} {candidate.get('sheet', '')} {candidate.get('missing_cells', '')}"
        ),
        "human_decision": "",
        "human_comment": "",
    }


def classify_pdf_candidate(candidate: dict[str, Any], locked_notes: set[int], final_by_note: dict[int, list[dict[str, Any]]]) -> dict[str, Any]:
    note = note_int(candidate.get("note")) or 0
    overlap = clean_text(candidate.get("current_gt_overlap"))
    current_hint, current_hint_detail = final_coverage_hint(candidate, final_by_note)
    if note in locked_notes:
        bucket = "locked_confirmed_skip"
        action = "skip_front11_confirmed"
        risk = "skip"
    elif overlap or current_hint:
        bucket = "pdf_value_maybe_covered"
        action = "no_add_first"
        risk = "low"
    else:
        bucket = "low_confidence_pdf_extra_value_pool"
        action = "do_not_review_one_by_one_until_excel_side_done"
        risk = "low"
    return {
        "candidate_id": clean_text(candidate.get("candidate_id")),
        "note": note,
        "folder": clean_text(candidate.get("folder")),
        "source": "pdf_to_excel",
        "bucket": bucket,
        "risk": risk,
        "recommended_action": action,
        "suggested_issue_type": "candidate:pdf_extra_value",
        "pdf_value": clean_text(candidate.get("pdf_value")),
        "pdf_value_key": clean_text(candidate.get("pdf_value_key")),
        "pdf_file": clean_text(candidate.get("pdf_file")),
        "xlsx_file": clean_text(candidate.get("xlsx_file")),
        "current_gt_overlap": overlap,
        "current_final_hint": current_hint_detail,
        "pdf_context": clean_text(candidate.get("pdf_context")),
        "review_instruction": "低置信池：PDF 单个数值不在 Excel，容易包含比较期/格式/相邻附注；不要逐条人工看，先按附注抽样。",
    }


def stable_candidate_key(row: dict[str, Any]) -> str:
    joined = "|".join(
        [
            clean_text(row.get("note")),
            clean_text(row.get("suggested_issue_type")),
            norm_text(row.get("table_path")),
            norm_text(row.get("section_path")),
            norm_text(row.get("row_label")),
            norm_text(row.get("column_label")),
            norm_text(row.get("period")),
            norm_text(row.get("source_value") or row.get("pdf_value")),
            norm_text(row.get("target_value")),
        ]
    )
    return hashlib.sha256(joined.encode("utf-8")).hexdigest()


def attach_candidate_keys(rows: list[dict[str, Any]]) -> None:
    for row in rows:
        row["candidate_key"] = stable_candidate_key(row)


def compact_raw_payload(row: dict[str, Any], context_chars: int) -> str:
    compact: dict[str, Any] = {}
    for key, value in row.items():
        if key in {"pdf_context", "review_instruction", "reason", "manual_rationale", "current_final_hint"}:
            compact[key] = clip_text(value, context_chars)
        else:
            compact[key] = value
    return clip_text(json.dumps(compact, ensure_ascii=False, separators=(",", ":")), max(4000, context_chars * 2))


def camel_candidate(row: dict[str, Any], context_chars: int = 2000) -> dict[str, Any]:
    pdf_only = row.get("source") == "pdf_to_excel"
    payload = {
        "candidateId": clean_text(row.get("candidate_id")),
        "candidateKey": clean_text(row.get("candidate_key")),
        "noteNo": clean_text(row.get("note")),
        "noteName": clean_text(row.get("note_name")),
        "folder": clean_text(row.get("folder")),
        "source": clean_text(row.get("source")),
        "bucket": clean_text(row.get("bucket")),
        "risk": clean_text(row.get("risk")),
        "recommendedAction": clean_text(row.get("recommended_action")),
        "suggestedIssueType": clean_text(row.get("suggested_issue_type")),
        "reason": clean_text(row.get("reason") or row.get("review_instruction")),
        "candidateKind": clean_text(row.get("candidate_kind")),
        "priority": row.get("priority"),
        "periodBucket": clean_text(row.get("period_bucket")),
        "tablePath": clean_text(row.get("table_path")),
        "sectionPath": clean_text(row.get("section_path")),
        "rowLabel": clean_text(row.get("row_label")),
        "columnLabel": clean_text(row.get("column_label")),
        "periodLabel": clean_text(row.get("period")),
        "sourceValue": clean_text(row.get("source_value") if not pdf_only else row.get("pdf_value")),
        "targetValue": clean_text(row.get("target_value")),
        "missingCells": clean_text(row.get("missing_cells")),
        "missingValues": clean_text(row.get("missing_values")),
        "currentGtOverlap": clean_text(row.get("current_gt_overlap")),
        "currentFinalHint": clean_text(row.get("current_final_hint")),
        "manualDecision": clean_text(row.get("manual_decision")),
        "manualRationale": clean_text(row.get("manual_rationale")),
        "sourceLocator": clean_text(row.get("source_locator") if not pdf_only else row.get("pdf_file")),
        "targetLocator": clean_text(row.get("target_locator") if not pdf_only else row.get("xlsx_file")),
        "pdfContext": clip_text(row.get("pdf_context"), context_chars),
        "rawPayloadJson": compact_raw_payload(row, context_chars),
    }
    return payload


def autosize(ws) -> None:
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 8
        for row_idx in range(1, min(ws.max_row, 200) + 1):
            max_len = max(max_len, len(clean_text(ws.cell(row_idx, col_idx).value)))
        ws.column_dimensions[letter].width = min(max_len + 2, 42)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def add_sheet(wb: Workbook, title: str, rows: list[dict[str, Any]], columns: list[str]) -> None:
    ws = wb.create_sheet(title)
    ws.append(columns)
    for row in rows:
        ws.append([row.get(col, "") for col in columns])
    header_fill = PatternFill("solid", fgColor="1F4E5F")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autosize(ws)


def write_csv(path: Path, rows: list[dict[str, Any]], columns: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_html(path: Path, summary: dict[str, Any], high_rows: list[dict[str, Any]], medium_rows: list[dict[str, Any]]) -> None:
    def esc(value: Any) -> str:
        return html.escape(clean_text(value))

    cards = []
    for row in (high_rows + medium_rows)[:80]:
        cards.append(
            f"""
            <article class="card {esc(row['risk'])}">
              <h3>{esc(row['candidate_id'])} · 附注 {esc(row['note'])} {esc(row['folder'])}</h3>
              <p><b>建议：</b>{esc(row['recommended_action'])} / {esc(row['suggested_issue_type'])}</p>
              <p><b>行名：</b>{esc(row['row_label'])}</p>
              <p><b>Excel 值：</b>{esc(row['target_value'])}</p>
              <p><b>原因：</b>{esc(row['reason'])}</p>
              <p><b>PDF 上下文：</b>{esc(row['pdf_context'])[:600]}</p>
              <p><b>Excel 定位：</b>{esc(row['target_locator'])}</p>
            </article>
            """
        )
    html_text = f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8" />
  <title>问题GT漏项候选审核包</title>
  <style>
    body {{ font-family: "Microsoft YaHei", Arial, sans-serif; margin: 28px; background: #f6f0df; color: #222; }}
    .summary {{ display:grid; grid-template-columns: repeat(4, minmax(0,1fr)); gap: 12px; margin: 18px 0; }}
    .metric {{ background:white; border-radius:16px; padding:16px; box-shadow: 0 6px 20px #0001; }}
    .metric b {{ display:block; font-size:28px; color:#1f4e5f; }}
    .card {{ background:white; border-left: 8px solid #999; border-radius:16px; padding:16px; margin:14px 0; box-shadow: 0 6px 20px #0001; }}
    .card.high {{ border-color:#b6402a; }}
    .card.medium {{ border-color:#c7922b; }}
    h1 {{ margin-bottom: 6px; }}
    p {{ line-height: 1.6; }}
  </style>
</head>
<body>
  <h1>问题GT漏项候选审核包 MVP</h1>
  <p>前 11 个附注已锁定，不进入本轮人工审核；本报告只辅助发现后续附注可能漏掉的 PDF vs Excel 问题GT。</p>
  <section class="summary">
    {''.join(f'<div class="metric"><span>{html.escape(k)}</span><b>{html.escape(str(v))}</b></div>' for k,v in summary.items())}
  </section>
  <h2>优先人工看的候选</h2>
  {''.join(cards) if cards else '<p>暂无高/中优先级候选。</p>'}
</body>
</html>"""
    path.write_text(html_text, encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="Build MVP review pack for problem GT missing candidates.")
    parser.add_argument("--final-gt", default=r"D:\data-annotation\final_gt_202506.xlsx")
    parser.add_argument("--source-candidates", default=r"D:\data-annotation\2025.06\gt_source_recheck_no_yellow_20260604\full_source_compare_no_yellow_candidates.json")
    parser.add_argument("--manual-decisions", default=r"D:\data-annotation\2025.06\gt_source_recheck_no_yellow_20260604\manual-reviewed-decisions-48-no-yellow-20260604.json")
    parser.add_argument("--output-dir", default="")
    parser.add_argument("--locked-notes", default="1-11")
    parser.add_argument("--pdf-sample-per-note", type=int, default=5)
    parser.add_argument("--source-run-key", default="")
    parser.add_argument("--backend", default="")
    parser.add_argument("--project-id", default="")
    parser.add_argument(
        "--backend-pdf-mode",
        choices=["none", "samples", "all"],
        default="samples",
        help="导入后端时 PDF->Excel 低置信池的范围：none=不导入，samples=每附注抽样，all=全量导入。",
    )
    parser.add_argument(
        "--xlsx-pdf-mode",
        choices=["summary", "samples", "all"],
        default="samples",
        help="生成 Excel 审核包时 PDF->Excel 低置信池的范围：summary=只按附注汇总，samples=每附注抽样，all=全量写入。",
    )
    parser.add_argument("--backend-context-chars", type=int, default=2000, help="导入后端时 PDF 上下文最长字符数。")
    args = parser.parse_args()

    final_gt = Path(args.final_gt)
    source_candidates = Path(args.source_candidates)
    manual_decisions_path = Path(args.manual_decisions) if args.manual_decisions else None
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    source_run_key = args.source_run_key or f"problem_gt_mvp_202506_{timestamp}"
    output_dir = Path(args.output_dir) if args.output_dir else Path(r"D:\data-annotation") / f"problem_gt_mvp_202506_{timestamp}"
    output_dir.mkdir(parents=True, exist_ok=True)

    locked_notes = parse_lock_notes(args.locked_notes)
    print(f"reading final_gt={final_gt}", flush=True)
    final_rows, final_by_note = read_final_gt(final_gt)
    print(f"reading source_candidates={source_candidates}", flush=True)
    source_data = read_json(source_candidates)
    print(f"reading manual_decisions={manual_decisions_path}", flush=True)
    manual_decisions = read_manual_decisions(manual_decisions_path)

    print("classifying Excel->PDF candidates", flush=True)
    excel_rows = [
        classify_excel_candidate(
            item,
            locked_notes=locked_notes,
            final_by_note=final_by_note,
            manual_decisions=manual_decisions,
        )
        for item in source_data.get("excel_to_pdf_candidates", [])
    ]
    print("classifying PDF->Excel candidates", flush=True)
    pdf_rows = [classify_pdf_candidate(item, locked_notes, final_by_note) for item in source_data.get("pdf_to_excel_candidates", [])]
    attach_candidate_keys(excel_rows)
    attach_candidate_keys(pdf_rows)

    high_rows = [r for r in excel_rows if r["bucket"] == "high_confidence_possible_missing"]
    medium_rows = [r for r in excel_rows if r["bucket"] in {"medium_confidence_possible_missing", "maybe_covered_by_current_final"}]
    manual_rows = [r for r in excel_rows if r["risk"] == "manual_reused"]
    covered_rows = [r for r in excel_rows if r["bucket"] in {"covered_by_source_gt_index", "maybe_covered_by_current_final"}]
    locked_rows = [r for r in excel_rows if r["bucket"] == "locked_confirmed_skip"]

    pdf_pool = [r for r in pdf_rows if r["bucket"] == "low_confidence_pdf_extra_value_pool"]
    pdf_by_note: list[dict[str, Any]] = []
    pdf_group: dict[int, list[dict[str, Any]]] = defaultdict(list)
    for row in pdf_pool:
        pdf_group[row["note"]].append(row)
    for note, rows in sorted(pdf_group.items()):
        pdf_by_note.append(
            {
                "note": note,
                "folder": rows[0].get("folder", ""),
                "candidate_count": len(rows),
                "sample_values": " | ".join(row["pdf_value"] for row in rows[: args.pdf_sample_per_note]),
                "recommended_action": "先不要逐条看；仅当 Excel 侧候选清完后，再按附注抽样确认是否存在 row:extra/text:extra。",
            }
        )
    pdf_samples = [row for rows in pdf_group.values() for row in rows[: args.pdf_sample_per_note]]
    if args.backend_pdf_mode == "all":
        backend_pdf_rows = pdf_rows
    elif args.backend_pdf_mode == "samples":
        backend_pdf_rows = pdf_samples
    else:
        backend_pdf_rows = []
    backend_rows = excel_rows + backend_pdf_rows
    if args.xlsx_pdf_mode == "all":
        xlsx_pdf_rows = pdf_rows
    elif args.xlsx_pdf_mode == "samples":
        xlsx_pdf_rows = pdf_samples
    else:
        xlsx_pdf_rows = []

    final_issue_counter = Counter(row["issue_type"] for row in final_rows)
    final_note_rows = []
    for note in sorted(final_by_note):
        rows = final_by_note[note]
        final_note_rows.append(
            {
                "note": note,
                "note_name": rows[0].get("note_name", ""),
                "locked": "yes" if note in locked_notes else "no",
                "final_gt_count": len(rows),
                "issue_types": "; ".join(f"{k}:{v}" for k, v in Counter(r["issue_type"] for r in rows).items()),
            }
        )
    locked_note_rows = [
        {
            "note": note,
            "note_name": LOCKED_NOTE_NAMES.get(note, ""),
            "status": "LOCKED_CONFIRMED",
            "note": note,
            "instruction": "用户确认该附注 100% 正确；MVP 不再生成审核任务。",
            "current_final_gt_count": len(final_by_note.get(note, [])),
        }
        for note in sorted(locked_notes)
    ]

    summary = {
        "final_gt_rows": len(final_rows),
        "locked_notes": len(locked_notes),
        "excel_candidates_total": len(excel_rows),
        "high_review_first": len(high_rows),
        "medium_review_after_high": len(medium_rows),
        "manual_reused": len(manual_rows),
        "covered_or_duplicate": len(covered_rows),
        "pdf_low_conf_pool": len(pdf_pool),
        "backend_import_candidates": len(backend_rows),
        "backend_pdf_mode": args.backend_pdf_mode,
        "xlsx_pdf_mode": args.xlsx_pdf_mode,
    }

    columns = [
        "candidate_id",
        "candidate_key",
        "note",
        "folder",
        "source",
        "bucket",
        "risk",
        "recommended_action",
        "suggested_issue_type",
        "reason",
        "candidate_kind",
        "priority",
        "period_bucket",
        "row_label",
        "column_label",
        "period",
        "source_value",
        "target_value",
        "missing_cells",
        "missing_values",
        "current_gt_overlap",
        "current_final_hint",
        "manual_decision",
        "manual_rationale",
        "source_locator",
        "target_locator",
        "pdf_context",
        "human_decision",
        "human_comment",
    ]
    pdf_columns = [
        "candidate_id",
        "candidate_key",
        "note",
        "folder",
        "source",
        "bucket",
        "risk",
        "recommended_action",
        "suggested_issue_type",
        "pdf_value",
        "pdf_file",
        "xlsx_file",
        "current_gt_overlap",
        "current_final_hint",
        "pdf_context",
        "review_instruction",
    ]

    wb = Workbook()
    wb.remove(wb.active)
    add_sheet(wb, "说明", [{"metric": k, "value": v} for k, v in summary.items()], ["metric", "value"])
    add_sheet(wb, "锁定前11附注", locked_note_rows, ["note", "note_name", "status", "current_final_gt_count", "instruction"])
    add_sheet(wb, "当前GT按附注", final_note_rows, ["note", "note_name", "locked", "final_gt_count", "issue_types"])
    add_sheet(wb, "优先审核_高置信", high_rows, columns)
    add_sheet(wb, "继续审核_中置信", medium_rows, columns)
    add_sheet(wb, "历史人工结论复用", manual_rows, columns)
    add_sheet(wb, "已有GT覆盖或疑似重复", covered_rows, columns)
    add_sheet(wb, "PDF低置信池_按附注", pdf_by_note, ["note", "folder", "candidate_count", "sample_values", "recommended_action"])
    add_sheet(wb, "PDF低置信池_样本", pdf_samples, pdf_columns)
    add_sheet(wb, "Excel候选全集", excel_rows, columns)
    if xlsx_pdf_rows:
        add_sheet(wb, "PDF候选后端范围_慎看", xlsx_pdf_rows, pdf_columns)
    xlsx_path = output_dir / "problem_gt_missing_review_pack.xlsx"
    wb.save(xlsx_path)

    write_csv(output_dir / "high_confidence_possible_missing.csv", high_rows, columns)
    write_csv(output_dir / "medium_confidence_possible_missing.csv", medium_rows, columns)
    write_csv(output_dir / "manual_decision_reuse.csv", manual_rows, columns)
    write_csv(output_dir / "pdf_low_confidence_by_note.csv", pdf_by_note, ["note", "folder", "candidate_count", "sample_values", "recommended_action"])
    write_html(output_dir / "problem_gt_missing_review_pack.html", summary, high_rows, medium_rows)
    import_payload = {
        "sourceRunKey": source_run_key,
        "finalGtPath": str(final_gt),
        "sourceCandidatesPath": str(source_candidates),
        "manualDecisionsPath": str(manual_decisions_path) if manual_decisions_path else "",
        "generatedAt": timestamp,
        "lockedNotes": [str(note) for note in sorted(locked_notes)],
        "candidates": [camel_candidate(row, args.backend_context_chars) for row in backend_rows],
    }
    payload_path = output_dir / "problem_gt_import_payload.json"
    payload_path.write_text(json.dumps(import_payload, ensure_ascii=False, indent=2), encoding="utf-8")

    manifest = {
        "created_at": timestamp,
        "source_run_key": source_run_key,
        "final_gt": str(final_gt),
        "source_candidates": str(source_candidates),
        "manual_decisions": str(manual_decisions_path) if manual_decisions_path else "",
        "locked_notes": sorted(locked_notes),
        "summary": summary,
        "final_issue_distribution": dict(final_issue_counter),
        "outputs": {
            "xlsx": str(xlsx_path),
            "html": str(output_dir / "problem_gt_missing_review_pack.html"),
            "high_csv": str(output_dir / "high_confidence_possible_missing.csv"),
            "medium_csv": str(output_dir / "medium_confidence_possible_missing.csv"),
            "backend_payload": str(payload_path),
        },
    }
    (output_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    if args.backend and args.project_id:
        url = f"{args.backend.rstrip('/')}/api/projects/{args.project_id}/problem-gt/import"
        result = post_json(url, import_payload)
        (output_dir / "backend_import_result.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"backend_import={url}")
        print(json.dumps(result, ensure_ascii=False, indent=2))

    summary_md = f"""# 问题GT漏项候选审核包 MVP

## 输入

- final GT: `{final_gt}`
- 历史候选: `{source_candidates}`
- 历史人工结论: `{manual_decisions_path}`
- 锁定附注: `{args.locked_notes}`（1 货币资金 到 11 融券业务）

## 核心数量

| 指标 | 数量 |
|---|---:|
| 当前 final GT 行数 | {summary['final_gt_rows']} |
| 锁定附注数 | {summary['locked_notes']} |
| Excel->PDF 候选总数 | {summary['excel_candidates_total']} |
| 高置信优先审核 | {summary['high_review_first']} |
| 中置信继续审核 | {summary['medium_review_after_high']} |
| 历史人工结论复用 | {summary['manual_reused']} |
| 已有 GT 覆盖/疑似重复 | {summary['covered_or_duplicate']} |
| PDF 数值低置信池 | {summary['pdf_low_conf_pool']} |
| 后端导入候选数 | {summary['backend_import_candidates']} |
| 后端 PDF 导入模式 | {summary['backend_pdf_mode']} |
| Excel 审核包 PDF 模式 | {summary['xlsx_pdf_mode']} |

## 使用顺序

1. 先看 `优先审核_高置信`，这些是最可能漏掉且数量最少的候选。
2. 再看 `继续审核_中置信`，重点判断是否已被当前 final GT 覆盖或是否只是辅助列/比较期。
3. `历史人工结论复用` 不需要重新看，除非你想抽查之前的人工判断。
4. `PDF低置信池_按附注` 不建议逐条看；等 Excel 侧候选清完后，再按附注抽样判断是否有 PDF extra 类漏项。

## 边界

- 本脚本没有修改 `final_gt_202506.xlsx`。
- 当前 MVP 优先减少人工量，不自动把候选写入最终 GT。
- PDF 数值候选噪声较高，不能直接当作 GT。
"""
    (output_dir / "summary.md").write_text(summary_md, encoding="utf-8")

    print(f"output_dir={output_dir}")
    print(f"xlsx={xlsx_path}")
    print(f"html={output_dir / 'problem_gt_missing_review_pack.html'}")
    print(f"backend_payload={payload_path}")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
