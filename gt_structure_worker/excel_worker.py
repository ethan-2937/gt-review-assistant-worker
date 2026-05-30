from __future__ import annotations

import argparse
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from gt_structure_worker.common import (
    CellItem,
    add_common_args,
    clean_text,
    first_non_blank,
    infer_note_from_heading,
    infer_note_from_path,
    infer_note_from_text,
    iter_files,
    matrix_to_table_payload,
    merge_note_payloads,
    non_empty_count,
    row_values,
    split_blocks,
    write_json,
)

EXCEL_SUFFIXES = (".xlsx", ".xlsm")
SKIP_SHEET_HINTS = ("主表", "目录")


def merged_value_lookup(ws) -> dict[tuple[int, int], Any]:
    lookup: dict[tuple[int, int], Any] = {}
    for merged_range in ws.merged_cells.ranges:
        top_left = ws.cell(merged_range.min_row, merged_range.min_col).value
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for col in range(merged_range.min_col, merged_range.max_col + 1):
                lookup[(row, col)] = top_left
    return lookup


def sheet_actual_bounds(ws, include_hidden: bool = False) -> tuple[int, int, int, int] | None:
    min_row = min_col = None
    max_row = max_col = None
    for row in ws.iter_rows():
        row_hidden = ws.row_dimensions[row[0].row].hidden if row else False
        if row_hidden and not include_hidden:
            continue
        for cell in row:
            col_hidden = ws.column_dimensions[get_column_letter(cell.column)].hidden
            if col_hidden and not include_hidden:
                continue
            if cell.value is None:
                continue
            if clean_text(cell.value) == "":
                continue
            min_row = cell.row if min_row is None else min(min_row, cell.row)
            max_row = cell.row if max_row is None else max(max_row, cell.row)
            min_col = cell.column if min_col is None else min(min_col, cell.column)
            max_col = cell.column if max_col is None else max(max_col, cell.column)
    if min_row is None:
        return None
    return min_row, max_row, min_col, max_col


def cell_item(ws, row: int, col: int, merged_lookup: dict[tuple[int, int], Any]) -> CellItem:
    value = ws.cell(row, col).value
    if value is None and (row, col) in merged_lookup:
        value = merged_lookup[(row, col)]
    locator = f"{Path(ws.parent.path).name if getattr(ws.parent, 'path', None) else 'workbook'}!{ws.title}!{get_column_letter(col)}{row}"
    return CellItem(value=value, locator=locator)


def sheet_to_matrix(ws, include_hidden: bool = False) -> list[list[CellItem]]:
    bounds = sheet_actual_bounds(ws, include_hidden=include_hidden)
    if not bounds:
        return []
    min_row, max_row, min_col, max_col = bounds
    merged_lookup = merged_value_lookup(ws)
    matrix: list[list[CellItem]] = []
    for row in range(min_row, max_row + 1):
        if ws.row_dimensions[row].hidden and not include_hidden:
            matrix.append([CellItem() for _ in range(min_col, max_col + 1)])
            continue
        matrix_row: list[CellItem] = []
        for col in range(min_col, max_col + 1):
            if ws.column_dimensions[get_column_letter(col)].hidden and not include_hidden:
                matrix_row.append(CellItem())
            else:
                matrix_row.append(cell_item(ws, row, col, merged_lookup))
        matrix.append(matrix_row)
    return matrix


def infer_note_for_sheet(file_path: Path, sheet_name: str, matrix: list[list[CellItem]]) -> tuple[str, str]:
    for text in (file_path.stem, file_path.parent.name, sheet_name):
        note_no, note_name = infer_note_from_text(text)
        if note_no:
            return note_no, note_name
    for row in matrix[:12]:
        row_text = " ".join(clean_text(cell.value) for cell in row if clean_text(cell.value))
        note_no, note_name = infer_note_from_text(row_text)
        if note_no:
            return note_no, note_name
    return infer_note_from_path(file_path)


def note_heading_from_row(row: list[CellItem]) -> tuple[str, str]:
    values = [value for value in row_values(row) if value]
    if not values or len(values) > 3:
        return "", ""
    row_text = " ".join(values)
    if len(row_text) > 80:
        return "", ""
    return infer_note_from_heading(row_text, allow_space_separator=True, strict=True)


def split_matrix_by_note_sections(
    file_path: Path,
    sheet_name: str,
    matrix: list[list[CellItem]],
) -> list[tuple[str, str, list[list[CellItem]], str]]:
    headings: list[tuple[int, str, str]] = []
    for idx, row in enumerate(matrix):
        note_no, note_name = note_heading_from_row(row)
        if note_no:
            headings.append((idx, note_no, note_name))

    sections: list[tuple[str, str, list[list[CellItem]], str]] = []
    if headings:
        for pos, (start_idx, note_no, note_name) in enumerate(headings):
            end_idx = headings[pos + 1][0] if pos + 1 < len(headings) else len(matrix)
            section = matrix[start_idx + 1:end_idx]
            if not any(non_empty_count(row) > 1 for row in section):
                continue
            locator = f"{file_path.name}!{sheet_name}!note{note_no}"
            sections.append((note_no, note_name, section, locator))
        return sections

    note_no, note_name = infer_note_for_sheet(file_path, sheet_name, matrix)
    if note_no:
        sections.append((note_no, note_name, matrix, f"{file_path.name}!{sheet_name}"))
    return sections


def parse_workbook(
    file_path: Path,
    *,
    include_hidden: bool = False,
    include_empty_cells: bool = False,
    blank_row_gap: int = 2,
) -> list[dict[str, Any]]:
    wb = load_workbook(file_path, data_only=True, read_only=False)
    # Save path on workbook object so locators can include the source file name.
    wb.path = str(file_path)

    note_tables: dict[str, dict[str, Any]] = {}
    table_order_by_note: dict[str, int] = {}

    for ws in wb.worksheets:
        if ws.sheet_state != "visible" and not include_hidden:
            continue
        if any(hint in ws.title for hint in SKIP_SHEET_HINTS):
            continue
        matrix = sheet_to_matrix(ws, include_hidden=include_hidden)
        if not matrix:
            continue
        sections = split_matrix_by_note_sections(file_path, ws.title, matrix)
        for note_no, note_name, section_matrix, section_locator in sections:
            if not note_no:
                continue
            note_key = note_no
            if note_key not in note_tables:
                note_tables[note_key] = {
                    "noteNo": note_no,
                    "noteName": first_non_blank(note_name, ""),
                    "sourceLocator": section_locator,
                    "tables": [],
                }
                table_order_by_note[note_key] = 1

            blocks = split_blocks(section_matrix, blank_gap=blank_row_gap)
            for block in blocks:
                order = table_order_by_note[note_key]
                table = matrix_to_table_payload(
                    block,
                    default_table_title=f"{ws.title} 附注{note_no} 表{order}",
                    table_order=order,
                    locator_prefix=f"{section_locator}!block{order}",
                    include_empty_cells=include_empty_cells,
                )
                if table:
                    note_tables[note_key]["tables"].append(table)
                    table_order_by_note[note_key] += 1

    wb.close()
    return list(note_tables.values())


def build_excel_payload(
    input_path: str | Path,
    *,
    include_hidden: bool = False,
    include_empty_cells: bool = False,
    blank_row_gap: int = 2,
) -> dict[str, Any]:
    files = iter_files(input_path, EXCEL_SUFFIXES)
    notes: list[dict[str, Any]] = []
    for file_path in files:
        notes.extend(parse_workbook(
            file_path,
            include_hidden=include_hidden,
            include_empty_cells=include_empty_cells,
            blank_row_gap=blank_row_gap,
        ))
    return {
        "side": "EXCEL",
        "sourceFilePath": str(Path(input_path)),
        "notes": merge_note_payloads(notes),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="把 Excel 财报附注拆成后端可导入的结构 JSON")
    add_common_args(parser)
    parser.add_argument("--include-hidden", action="store_true", help="包含隐藏 sheet/行/列")
    parser.add_argument("--blank-row-gap", type=int, default=2, help="连续多少个空行视为拆表边界，默认 2")
    args = parser.parse_args()

    payload = build_excel_payload(
        args.input,
        include_hidden=args.include_hidden,
        include_empty_cells=args.include_empty_cells,
        blank_row_gap=args.blank_row_gap,
    )
    write_json(payload, args.output)
    print(f"EXCEL notes={len(payload['notes'])} output={args.output}")


if __name__ == "__main__":
    main()
