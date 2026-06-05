from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from urllib import parse as urlparse
from urllib import request as urlrequest

import fitz
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from PIL import Image, ImageDraw, ImageFont


DEFAULT_SOURCE_ROOT = r"D:\data-annotation\2025.06\2025上半年-附注分割版-标黄-仅2025-复核"
DEFAULT_OUTPUT_DIR = r"D:\audit-engine\gt-review-assistant\workspace\preview_assets"


@dataclass(frozen=True)
class CandidatePreview:
    candidate_id: str
    candidate_key: str
    note_no: str
    note_name: str
    excel_locator: str
    pdf_locator: str
    excel_png: str
    pdf_png: str
    excel_status: str
    pdf_status: str


def clean_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def safe_part(value: Any) -> str:
    cleaned = clean_text(value)
    cleaned = re.sub(r"[^A-Za-z0-9_.-]+", "_", cleaned).strip("_")
    if cleaned:
        return cleaned
    return hashlib.md5(clean_text(value).encode("utf-8")).hexdigest()


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def download_proposal(backend: str, project_id: int, source_run_key: str, output_dir: Path) -> Path:
    backend = backend.rstrip("/")
    query = ""
    if clean_text(source_run_key):
        query = "?" + urlparse.urlencode({"sourceRunKey": source_run_key})
    url = f"{backend}/api/projects/{project_id}/problem-gt/export.json{query}"
    req = urlrequest.Request(url, headers={"Accept": "application/json"}, method="GET")
    with urlrequest.urlopen(req, timeout=300) as response:
        body = response.read()
    output_dir.mkdir(parents=True, exist_ok=True)
    proposal_path = output_dir / f"problem_gt_proposal_{safe_part(source_run_key)}.json"
    proposal_path.write_bytes(body)
    return proposal_path


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        Path("C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/simsun.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc" if bold else "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/truetype/noto/NotoSansCJK-Bold.ttc" if bold else "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSerifCJK-Bold.ttc" if bold else "/usr/share/fonts/opentype/noto/NotoSerifCJK-Regular.ttc"),
    ]
    for candidate in candidates:
        if candidate.exists():
            try:
                return ImageFont.truetype(str(candidate), size=size)
            except Exception:
                continue
    return ImageFont.load_default()


def text_color_for(bg: tuple[int, int, int]) -> tuple[int, int, int]:
    brightness = (bg[0] * 299 + bg[1] * 587 + bg[2] * 114) / 1000
    return (30, 30, 30) if brightness > 145 else (255, 255, 255)


def cell_fill(cell) -> tuple[int, int, int]:
    color = getattr(cell.fill.fgColor, "rgb", None)
    if color and isinstance(color, str) and len(color) in (6, 8):
        rgb = color[-6:]
        try:
            return tuple(int(rgb[idx:idx + 2], 16) for idx in (0, 2, 4))
        except ValueError:
            pass
    return 255, 255, 255


def scan_files(source_roots: list[Path], suffix: str) -> list[Path]:
    files: list[Path] = []
    for root in source_roots:
        if root.is_file() and root.suffix.lower() == suffix:
            files.append(root)
        elif root.exists():
            files.extend(root.rglob(f"*{suffix}"))
    return files


def build_file_index(source_roots: list[Path]) -> dict[str, list[Path]]:
    index: dict[str, list[Path]] = {"xlsx": [], "pdf": []}
    index["xlsx"] = scan_files(source_roots, ".xlsx")
    index["pdf"] = scan_files(source_roots, ".pdf")
    return index


def note_tokens(note_no: str, note_name: str) -> list[str]:
    tokens = []
    if clean_text(note_no):
        tokens.append(clean_text(note_no))
    name = clean_text(note_name)
    if name:
        tokens.append(name)
        if "-" in name:
            tokens.extend(part.strip() for part in name.split("-") if part.strip())
    return tokens


def locator_filename(locator: str, suffix: str) -> str:
    pattern = rf"([^;|]+?\.{re.escape(suffix.lstrip('.'))})"
    match = re.search(pattern, locator, flags=re.IGNORECASE)
    return clean_text(match.group(1)) if match else ""


def find_source_file(files: list[Path], locator: str, note_no: str, note_name: str, suffix: str) -> Path | None:
    named = locator_filename(locator, suffix)
    if named:
        named_key = named.replace("\\", "/").split("/")[-1].lower()
        for file in files:
            if file.name.lower() == named_key:
                return file

    tokens = note_tokens(note_no, note_name)
    note = clean_text(note_no)
    candidates = []
    for file in files:
        name = file.name.lower()
        parent = file.parent.name.lower()
        score = 0
        if note and re.match(rf"^{re.escape(note)}(\D|$)", file.stem):
            score += 5
        for token in tokens:
            token_lower = token.lower()
            if token_lower and token_lower in name:
                score += 3
            if token_lower and token_lower in parent:
                score += 2
        if score:
            candidates.append((score, len(str(file)), file))
    if not candidates:
        return None
    candidates.sort(key=lambda item: (-item[0], item[1]))
    return candidates[0][2]


def parse_cells(locator: str) -> list[str]:
    cells = [match.group(1).upper() for match in re.finditer(r"\b([A-Za-z]{1,3}\d{1,7})\b", locator)]
    seen: set[str] = set()
    result: list[str] = []
    for cell in cells:
        if cell not in seen:
            seen.add(cell)
            result.append(cell)
    return result


def parse_sheet_name(locator: str, note_no: str, workbook_sheets: list[str]) -> str:
    for pattern in [r"sheet_name\s*=\s*([^;,\s]+)", r"\b(Sheet[^\s;|,]*)\b"]:
        match = re.search(pattern, locator, flags=re.IGNORECASE)
        if match:
            candidate = clean_text(match.group(1))
            for sheet in workbook_sheets:
                if sheet.lower() == candidate.lower():
                    return sheet
    if note_no:
        expected = f"Sheet{note_no}"
        for sheet in workbook_sheets:
            if sheet.lower() == expected.lower():
                return sheet
    return workbook_sheets[0]


def cell_to_rc(cell_ref: str) -> tuple[int, int] | None:
    match = re.fullmatch(r"([A-Z]+)(\d+)", cell_ref.upper())
    if not match:
        return None
    return int(match.group(2)), column_index_from_string(match.group(1))


def norm_search(value: Any) -> str:
    return re.sub(r"[\s,，、/\\_：:;；()（）\[\]【】<>《》-]+", "", clean_text(value)).lower()


def number_tokens(value: Any) -> list[str]:
    tokens = []
    for token in re.findall(r"-?\d+(?:,\d{3})*(?:\.\d+)?|-?\d+(?:\.\d+)?", clean_text(value)):
        normalized = token.replace(",", "")
        if len(normalized.replace(".", "").replace("-", "")) >= 2:
            tokens.append(normalized)
    return tokens


def first_non_empty_col(ws, row_idx: int) -> int:
    for col_idx in range(1, min(ws.max_column, 20) + 1):
        if clean_text(ws.cell(row_idx, col_idx).value):
            return col_idx
    return 1


def find_excel_anchor(ws, row_data: dict[str, Any]) -> tuple[int, int, set[tuple[int, int]], str] | None:
    row_label = norm_search(row_data.get("row_label"))
    column_label = norm_search(row_data.get("column_label"))
    value_numbers = set(number_tokens(row_data.get("excel_value")) + number_tokens(row_data.get("missing_values")))
    best: tuple[int, int, int, set[tuple[int, int]], str] | None = None
    max_row = min(ws.max_row, 500)
    max_col = min(ws.max_column, 80)
    for row_idx in range(1, max_row + 1):
        row_values = [clean_text(ws.cell(row_idx, col_idx).value) for col_idx in range(1, max_col + 1)]
        row_text = " ".join(row_values)
        row_norm = norm_search(row_text)
        score = 0
        highlight: set[tuple[int, int]] = set()
        if row_label and row_label in row_norm:
            score += 20
        if column_label and column_label in row_norm:
            score += 4
        for col_idx, value in enumerate(row_values, start=1):
            cell_numbers = set(number_tokens(value))
            if value_numbers and cell_numbers.intersection(value_numbers):
                score += min(12, len(cell_numbers.intersection(value_numbers)) * 4)
                highlight.add((row_idx, col_idx))
        if score <= 0:
            continue
        anchor_col = min((col for _, col in highlight), default=first_non_empty_col(ws, row_idx))
        status = "ok_fuzzy_row" if row_label and row_label in row_norm else "ok_fuzzy_value"
        candidate = (score, row_idx, anchor_col, highlight, status)
        if best is None or candidate[0] > best[0]:
            best = candidate
    if best is None:
        return None
    _, row_idx, col_idx, highlight, status = best
    return row_idx, col_idx, highlight, status


def render_excel_preview(excel_path: Path, locator: str, note_no: str, output_path: Path, row_data: dict[str, Any]) -> str:
    cells = parse_cells(locator)
    wb = load_workbook(excel_path, data_only=True, read_only=False)
    try:
        sheet_name = parse_sheet_name(locator, note_no, wb.sheetnames)
        ws = wb[sheet_name]
        status = "ok"
        if cells:
            first = cell_to_rc(cells[0])
            if not first:
                return "invalid_excel_cell"
            target_row, target_col = first
            highlight = {cell_to_rc(cell) for cell in cells}
            highlight = {item for item in highlight if item}
        else:
            anchor = find_excel_anchor(ws, row_data)
            if not anchor:
                for fallback_sheet in wb.sheetnames:
                    if fallback_sheet == sheet_name:
                        continue
                    fallback_ws = wb[fallback_sheet]
                    fallback_anchor = find_excel_anchor(fallback_ws, row_data)
                    if fallback_anchor:
                        sheet_name = fallback_sheet
                        ws = fallback_ws
                        anchor = fallback_anchor
                        break
            if not anchor:
                return "no_excel_cell_or_fuzzy_match"
            target_row, target_col, highlight, status = anchor

        min_row = max(1, target_row - 5)
        max_row = min(ws.max_row, target_row + 6)
        min_col = max(1, target_col - 4)
        max_col = min(ws.max_column, target_col + 6)
        row_h = 34
        col_w = 150
        row_head_w = 56
        title_h = 56
        width = row_head_w + (max_col - min_col + 1) * col_w + 2
        height = title_h + row_h + (max_row - min_row + 1) * row_h + 2
        image = Image.new("RGB", (width, height), (250, 248, 241))
        draw = ImageDraw.Draw(image)
        normal_font = font(14)
        small_font = font(12)
        title_font = font(18, bold=True)

        draw.rectangle([0, 0, width, title_h], fill=(36, 79, 63))
        title = f"{excel_path.name} / {sheet_name} / {'; '.join(cells[:4])}"
        draw.text((14, 14), title[:90], fill=(255, 250, 240), font=title_font)

        for col in range(min_col, max_col + 1):
            x = row_head_w + (col - min_col) * col_w
            draw.rectangle([x, title_h, x + col_w, title_h + row_h], fill=(231, 241, 235), outline=(193, 203, 194))
            draw.text((x + 8, title_h + 8), get_column_letter(col), fill=(48, 54, 47), font=small_font)

        for row in range(min_row, max_row + 1):
            y = title_h + row_h + (row - min_row) * row_h
            draw.rectangle([0, y, row_head_w, y + row_h], fill=(231, 241, 235), outline=(193, 203, 194))
            draw.text((8, y + 8), str(row), fill=(48, 54, 47), font=small_font)
            for col in range(min_col, max_col + 1):
                x = row_head_w + (col - min_col) * col_w
                cell = ws.cell(row, col)
                bg = cell_fill(cell)
                if bg == (0, 0, 0):
                    bg = (255, 255, 255)
                draw.rectangle([x, y, x + col_w, y + row_h], fill=bg, outline=(213, 213, 205))
                value = clean_text(cell.value)
                if len(value) > 18:
                    value = value[:18] + "..."
                draw.text((x + 7, y + 8), value, fill=text_color_for(bg), font=normal_font)
                if (row, col) in highlight:
                    draw.rectangle([x + 2, y + 2, x + col_w - 2, y + row_h - 2], outline=(196, 85, 45), width=4)
                elif not highlight and row == target_row:
                    draw.rectangle([x + 2, y + 2, x + col_w - 2, y + row_h - 2], outline=(244, 180, 0), width=2)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        image.save(output_path)
        return status
    finally:
        wb.close()


def parse_pdf_page(locator: str) -> int:
    for pattern in [
        r"(?:^|!)p(\d{1,5})(?:!|$)",
        r"page\s*[=:]\s*(\d{1,5})",
        r"第\s*(\d{1,5})\s*页",
    ]:
        match = re.search(pattern, locator, flags=re.IGNORECASE)
        if match:
            return max(int(match.group(1)), 1)
    return 1


def pdf_search_terms(row_data: dict[str, Any]) -> list[str]:
    terms: list[str] = []
    for key in ["row_label", "column_label"]:
        value = clean_text(row_data.get(key))
        if len(value) >= 2:
            terms.append(value)
    for key in ["pdf_value", "excel_value"]:
        value = clean_text(row_data.get(key))
        if value and value.lower() != "empty":
            terms.extend(number_tokens(value)[:3])
            if len(value) <= 40:
                terms.append(value)
    seen: set[str] = set()
    result: list[str] = []
    for term in terms:
        term = clean_text(term)
        if term and term not in seen:
            seen.add(term)
            result.append(term)
    return result


def pdf_clip_for_terms(page, row_data: dict[str, Any]) -> tuple[Any, str]:
    for term in pdf_search_terms(row_data):
        rects = page.search_for(term)
        if not rects and "," in term:
            rects = page.search_for(term.replace(",", ""))
        if not rects:
            continue
        rect = rects[0]
        clip = fitz.Rect(
            0,
            max(0, rect.y0 - 100),
            page.rect.width,
            min(page.rect.height, rect.y1 + 180),
        )
        return clip, "ok_crop"
    return page.rect, "ok_page"


def pdf_clip_for_doc(doc, preferred_page_index: int, row_data: dict[str, Any]) -> tuple[int, Any, str]:
    page = doc.load_page(preferred_page_index)
    clip, status = pdf_clip_for_terms(page, row_data)
    if status == "ok_crop":
        return preferred_page_index, clip, status

    # Missing-row PDF locators are often empty. Search the split note PDF for
    # nearby row/value/column text so reviewers still see the closest page.
    for page_index in range(doc.page_count):
        if page_index == preferred_page_index:
            continue
        page = doc.load_page(page_index)
        clip, status = pdf_clip_for_terms(page, row_data)
        if status == "ok_crop":
            return page_index, clip, "ok_crop_search"
    return preferred_page_index, doc.load_page(preferred_page_index).rect, "ok_page"


def render_pdf_preview(pdf_path: Path, locator: str, output_path: Path, row_data: dict[str, Any]) -> str:
    page_no = parse_pdf_page(locator)
    doc = fitz.open(pdf_path)
    try:
        if doc.page_count <= 0:
            return "empty_pdf"
        page_index = min(max(page_no - 1, 0), doc.page_count - 1)
        page_index, clip, status = pdf_clip_for_doc(doc, page_index, row_data)
        page = doc.load_page(page_index)
        zoom = 1.8 if status.startswith("ok_crop") else 1.2
        pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), clip=clip, alpha=False)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        pix.save(str(output_path))
        return status
    finally:
        doc.close()


def write_csv(path: Path, rows: list[CandidatePreview]) -> None:
    headers = list(CandidatePreview.__dataclass_fields__.keys())
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            writer.writerow(row.__dict__)


def generate_previews(
    proposal: dict[str, Any],
    source_roots: list[Path],
    output_dir: Path,
    limit: int,
    include_buckets: set[str],
) -> dict[str, Any]:
    source_run_key = clean_text(proposal.get("source_run_key")) or "latest"
    run_part = safe_part(source_run_key)
    preview_root = output_dir / "problem_gt" / run_part
    files = build_file_index(source_roots)
    rows = proposal.get("all_candidates") or []
    if include_buckets:
        rows = [row for row in rows if clean_text(row.get("bucket")) in include_buckets]
    if limit > 0:
        rows = rows[:limit]

    results: list[CandidatePreview] = []
    stats = {
        "source_run_key": source_run_key,
        "candidate_count": len(rows),
        "excel_ok": 0,
        "pdf_ok": 0,
        "excel_missing": 0,
        "pdf_missing": 0,
    }
    for row in rows:
        candidate_key = clean_text(row.get("candidate_key")) or clean_text(row.get("candidate_id"))
        candidate_id = clean_text(row.get("candidate_id"))
        note_no = clean_text(row.get("note_no"))
        note_name = clean_text(row.get("note_name"))
        excel_locator = clean_text(row.get("excel_locator"))
        pdf_locator = clean_text(row.get("pdf_locator"))
        item_part = safe_part(candidate_key)
        excel_png = preview_root / f"{item_part}_excel.png"
        pdf_png = preview_root / f"{item_part}_pdf.png"

        excel_file = find_source_file(files["xlsx"], excel_locator, note_no, note_name, ".xlsx")
        if excel_file:
            excel_status = render_excel_preview(excel_file, excel_locator, note_no, excel_png, row)
        else:
            excel_status = "excel_file_not_found"
        if excel_status.startswith("ok"):
            stats["excel_ok"] += 1
        else:
            stats["excel_missing"] += 1

        pdf_file = find_source_file(files["pdf"], pdf_locator, note_no, note_name, ".pdf")
        if pdf_file:
            pdf_status = render_pdf_preview(pdf_file, pdf_locator, pdf_png, row)
        else:
            pdf_status = "pdf_file_not_found"
        if pdf_status.startswith("ok"):
            stats["pdf_ok"] += 1
        else:
            stats["pdf_missing"] += 1

        results.append(
            CandidatePreview(
                candidate_id=candidate_id,
                candidate_key=candidate_key,
                note_no=note_no,
                note_name=note_name,
                excel_locator=excel_locator,
                pdf_locator=pdf_locator,
                excel_png=str(excel_png),
                pdf_png=str(pdf_png),
                excel_status=excel_status,
                pdf_status=pdf_status,
            )
        )

    preview_root.mkdir(parents=True, exist_ok=True)
    write_csv(preview_root / "problem_gt_preview_manifest.csv", results)
    (preview_root / "problem_gt_preview_summary.json").write_text(json.dumps(stats, ensure_ascii=False, indent=2), encoding="utf-8")
    return stats


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate PNG previews for problem-GT PDF/Excel source locations.")
    parser.add_argument("--proposal-json", default="", help="Local proposal JSON. If omitted, use --backend and --project-id.")
    parser.add_argument("--backend", default="", help="Backend URL, for example http://localhost:8080")
    parser.add_argument("--project-id", type=int, default=0)
    parser.add_argument("--source-run-key", default="")
    parser.add_argument("--source-root", action="append", default=[], help="Source root containing split note PDF/XLSX files.")
    parser.add_argument("--output-dir", default=DEFAULT_OUTPUT_DIR, help="Usually workspace/preview_assets")
    parser.add_argument("--limit", type=int, default=0, help="0 means all candidates.")
    parser.add_argument("--include-bucket", action="append", default=[], help="Optional bucket filter; can be repeated.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    output_dir = Path(args.output_dir)
    if args.proposal_json:
        proposal_path = Path(args.proposal_json)
    else:
        if not args.backend or not args.project_id:
            raise RuntimeError("either --proposal-json or both --backend and --project-id are required")
        proposal_path = download_proposal(args.backend, args.project_id, args.source_run_key, output_dir)
    if not proposal_path.exists():
        raise FileNotFoundError(proposal_path)

    roots = [Path(root) for root in args.source_root] or [Path(DEFAULT_SOURCE_ROOT)]
    proposal = read_json(proposal_path)
    stats = generate_previews(
        proposal=proposal,
        source_roots=roots,
        output_dir=output_dir,
        limit=args.limit,
        include_buckets={clean_text(bucket) for bucket in args.include_bucket if clean_text(bucket)},
    )
    print(json.dumps(stats, ensure_ascii=False))


if __name__ == "__main__":
    main()
